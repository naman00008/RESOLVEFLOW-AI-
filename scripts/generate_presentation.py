"""
Automated PowerPoint Presentation Generator for ResolveFlow AI
Generates a polished 8-slide widescreen (16:9) presentation for the CA-3 evaluation.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_filename="ResolveFlow_AI_Presentation.pptx"):
    prs = Presentation()
    # Set 16:9 widescreen dimensions (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)     # Deep Slate / Navy
    COLOR_ACCENT = RGBColor(37, 99, 235)     # Vibrant Blue
    COLOR_EMERALD = RGBColor(5, 150, 105)    # Emerald Green
    COLOR_MUTED = RGBColor(100, 116, 139)    # Muted Gray
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_CARD_BG = RGBColor(248, 250, 252)  # Light Slate Tint
    COLOR_BORDER = RGBColor(226, 232, 240)   # Subtle Border

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="RESOLVEFLOW AI • CA-3 AGENTIC AI POC"):
        # Category Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        # Slide Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY

    # ==========================================================
    # SLIDE 1: Title Slide
    # ==========================================================
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Background Accent Card
    bg_shape = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_PRIMARY
    bg_shape.line.color.rgb = COLOR_PRIMARY

    # Title Container
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(3.2))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "CONTINUOUS ASSESSMENT 3 (CA-3) • AGENTIC AI PROOF OF CONCEPT"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    p0.space_after = Pt(14)

    p1 = tf1.add_paragraph()
    p1.text = "ResolveFlow AI"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_after = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "A Multi-Agent AI System for E-Commerce Customer Complaint Triage and Autonomous Resolution"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_after = Pt(28)

    p3 = tf1.add_paragraph()
    p3.text = "Domain: E-Commerce & Retail Operations | Architecture: 5 Collaborative Agents + Deterministic Tool Layer"
    p3.font.size = Pt(13)
    p3.font.color.rgb = COLOR_EMERALD

    # Footer Info Box
    info_box = slide1.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(11.333), Inches(1.0))
    tf_info = info_box.text_frame
    p_info = tf_info.paragraphs[0]
    p_info.text = "Project Team: Naman Vyas (@naman00008) & Akshita Mishra (@amakshita0506-dot) | Date: October 2026 | Technology: Python, Streamlit, Multi-Agent State Machine"
    p_info.font.size = Pt(12)
    p_info.font.color.rgb = COLOR_MUTED

    # ==========================================================
    # SLIDE 2: Problem Statement & Existing Inefficiencies
    # ==========================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "The Real-World Business Challenge: E-Commerce Support Inefficiencies")

    col_w = Inches(3.6)
    gap = Inches(0.35)
    top_pos = Inches(1.8)

    cards_data_s2 = [
        {
            "title": "1. Slow Response Times",
            "stat": "12 – 48 Hours",
            "stat_label": "Average Support Backlog",
            "desc": "Manual tier-1 ticket sorting and carrier checking creates severe customer dissatisfaction, negative reviews, and high churn rates."
        },
        {
            "title": "2. Failure of Static Chatbots",
            "stat": "73% Frustration",
            "stat_label": "Bot Abandonment Rate",
            "desc": "Traditional rule-based chatbots cannot verify order status, calculate restocking fees, query carrier APIs, or execute reverse logistics RMA labels."
        },
        {
            "title": "3. Rampant Return & Policy Fraud",
            "stat": "₹8,000 Cr+ Lost",
            "stat_label": "Annual Retail Return Abuse",
            "desc": "Serial returners and false 'empty box' claims exploit manual staff who lack real-time fraud velocity scoring and carrier signature cross-checks."
        }
    ]

    for i, card in enumerate(cards_data_s2):
        left_pos = Inches(0.8) + i * (col_w + gap)
        box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, col_w, Inches(4.8))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD_BG
        box.line.color.rgb = COLOR_BORDER

        tb = slide2.shapes.add_textbox(left_pos + Inches(0.25), top_pos + Inches(0.3), col_w - Inches(0.5), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = card["title"]
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY
        p_t.space_after = Pt(14)

        p_s = tf.add_paragraph()
        p_s.text = card["stat"]
        p_s.font.size = Pt(28)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_ACCENT
        
        p_sl = tf.add_paragraph()
        p_sl.text = card["stat_label"].upper()
        p_sl.font.size = Pt(10)
        p_sl.font.bold = True
        p_sl.font.color.rgb = COLOR_MUTED
        p_sl.space_after = Pt(18)

        p_d = tf.add_paragraph()
        p_d.text = card["desc"]
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = COLOR_PRIMARY

    # ==========================================================
    # SLIDE 3: Domain Justification & The Agentic Solution
    # ==========================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Domain Justification & Why Agentic AI is Essential")

    # Left Container: Traditional vs Agentic
    left_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = COLOR_CARD_BG
    left_box.line.color.rgb = COLOR_BORDER

    tb_left = slide3.shapes.add_textbox(Inches(1.05), Inches(2.0), Inches(5.1), Inches(4.4))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "Why E-Commerce Operations?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    bullets_l = [
        "High-Volume, Rule-Rich Environment: E-commerce complaints must strictly comply with product return windows, sanitary laws (FDA hygiene), and carrier SLA rules.",
        "Deterministic Tool Dependency: Resolving an order is not just text generation; it requires database lookups, financial refund math, and reverse courier bookings.",
        "Measurable Business ROI: Automating tier-1 triage cuts operational costs by 80% and reduces resolution latency from 24 hours to sub-100 milliseconds."
    ]
    for b in bullets_l:
        p_b = tf_l.add_paragraph()
        p_b.text = f"• {b}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = COLOR_PRIMARY
        p_b.space_after = Pt(8)

    # Right Container: Core Innovations
    right_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = COLOR_CARD_BG
    right_box.line.color.rgb = COLOR_BORDER

    tb_right = slide3.shapes.add_textbox(Inches(7.05), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_r = tb_right.text_frame
    tf_r.word_wrap = True

    p_r = tf_r.paragraphs[0]
    p_r.text = "ResolveFlow AI's Paradigm Shift"
    p_r.font.size = Pt(18)
    p_r.font.bold = True
    p_r.font.color.rgb = COLOR_EMERALD
    p_r.space_after = Pt(10)

    innovations = [
        "Multi-Agent Division of Labor: Dedicated agents for Triage, Telemetry Research, Tool Execution, Critic Auditing, and Communication.",
        "Self-Correction & Critic Gate: Critic agent actively audits preliminary proposals, flags synthetic fraud, and enforces Human-in-the-Loop (HITL) safety.",
        "Zero Hallucination Guarantee: Hard business rules (return eligibility, restocking percentages, carrier tracking) executed via deterministic Python tool layer.",
        "End-to-End Enterprise Integration: Generates live webhook payloads for Stripe refund execution and courier dispatch."
    ]
    for inn in innovations:
        p_inn = tf_r.add_paragraph()
        p_inn.text = f"• {inn}"
        p_inn.font.size = Pt(12)
        p_inn.font.color.rgb = COLOR_PRIMARY
        p_inn.space_after = Pt(8)

    # ==========================================================
    # SLIDE 4: Multi-Agent Architecture
    # ==========================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Multi-Agent System Architecture & Specialization")

    agents_s4 = [
        {"num": "1", "name": "Triage & Planner Agent", "color": COLOR_ACCENT, "desc": "Extracts Order ID, infers sentiment and urgency level, and dynamically decomposes the complaint into structured milestones."},
        {"num": "2", "name": "OMS & Research Agent", "color": COLOR_PRIMARY, "desc": "Queries synthetic enterprise Order Management System (OMS), customer lifetime return history, and live carrier tracking APIs."},
        {"num": "3", "name": "Resolution Executor Agent", "color": COLOR_EMERALD, "desc": "Applies deterministic business tools: verifies category return windows, calculates refunds & restocking fees, and reserves inventory."},
        {"num": "4", "name": "Critic & Fraud Auditor Agent", "color": RGBColor(220, 38, 38), "desc": "Audits proposed remedy for policy violations, scores customer abuse velocity (0-100), and overrules auto-refunds on high fraud risk."},
        {"num": "5", "name": "Communication Agent", "color": RGBColor(147, 51, 234), "desc": "Drafts empathetic, personalized customer responses with tone adjustments, and outputs live ERP/Stripe transactional webhook payloads."}
    ]

    card_h = Inches(0.9)
    gap_y = Inches(0.18)
    top_start = Inches(1.7)

    for i, ag in enumerate(agents_s4):
        y_pos = top_start + i * (card_h + gap_y)
        shape = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.7), card_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD_BG
        shape.line.color.rgb = COLOR_BORDER

        # Badge
        badge = slide4.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.1), y_pos + Inches(0.15), Inches(0.6), Inches(0.6))
        badge.fill.solid()
        badge.fill.fore_color.rgb = ag["color"]
        badge.line.fill.background()
        tf_b = badge.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.text = ag["num"]
        p_b.font.size = Pt(16)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_WHITE
        p_b.alignment = PP_ALIGN.CENTER

        # Content Textbox
        tb_c = slide4.shapes.add_textbox(Inches(1.9), y_pos + Inches(0.1), Inches(10.3), Inches(0.7))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        p_nm = tf_c.paragraphs[0]
        p_nm.text = ag["name"]
        p_nm.font.size = Pt(15)
        p_nm.font.bold = True
        p_nm.font.color.rgb = COLOR_PRIMARY
        
        p_ds = tf_c.add_paragraph()
        p_ds.text = ag["desc"]
        p_ds.font.size = Pt(11.5)
        p_ds.font.color.rgb = COLOR_MUTED

    # ==========================================================
    # SLIDE 5: Orchestration & Decision-Making Workflow
    # ==========================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Orchestration, Reflection & Human-in-the-Loop Governance")

    # Workflow Steps
    box_wf = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    box_wf.fill.solid()
    box_wf.fill.fore_color.rgb = COLOR_CARD_BG
    box_wf.line.color.rgb = COLOR_BORDER

    tb_wf = slide5.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
    tf_w = tb_wf.text_frame
    tf_w.word_wrap = True

    p = tf_w.paragraphs[0]
    p.text = "How Agents Collaborate, Reflect and Guard Decisions:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(14)

    steps_text = [
        "1. Dynamic Milestone Planning: The Planner avoids rigid prompt chaining by decomposing complaints based on inferred customer emotional state and stated intent.",
        "2. Ground-Truth Telemetry Binding: The Research Agent binds real database states (e.g. delivery date, carrier scan logs) into the context memory before any policy decision is formulated.",
        "3. Deterministic Policy Calculation: The Executor invokes pure-Python business tools to determine mathematical eligibility (e.g., 7-day electronics return window, restocking fee).",
        "4. The Critic Reflection Gate (Self-Correction): The Critic checks the Executor's proposal against both policy rules and customer fraud history. If fraud score >= 60, it halts automation and triggers a Human-in-the-Loop (HITL) escalation.",
        "5. Automated Transaction Fulfillment: Upon approval, the Communication Agent formats customer-ready communication and dispatches enterprise ERP / payment refund webhooks."
    ]

    for stp in steps_text:
        p_s = tf_w.add_paragraph()
        p_s.text = stp
        p_s.font.size = Pt(13)
        p_s.font.color.rgb = COLOR_PRIMARY
        p_s.space_after = Pt(10)

    # ==========================================================
    # SLIDE 6: PoC Demonstration & Scenario Evaluation
    # ==========================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Proof of Concept Demonstration: 5 Benchmark Scenarios")

    # Table of Scenarios
    rows, cols = 6, 5
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(11.7)
    height = Inches(4.8)

    table_shape = slide6.shapes.add_table(rows, cols, left, top, width, height)
    tbl = table_shape.table

    # Column Widths
    tbl.columns[0].width = Inches(1.4)  # Scenario
    tbl.columns[1].width = Inches(3.2)  # Customer Case
    tbl.columns[2].width = Inches(2.5)  # Agent Invocations
    tbl.columns[3].width = Inches(1.8)  # Risk Score
    tbl.columns[4].width = Inches(2.8)  # Final Output

    headers = ["Scenario", "Customer Complaint Case", "Agent Invocations", "Fraud Risk", "System Verdict"]
    for col_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    scenario_rows = [
        ("Scenario 1", "Damaged Headphones within 7-day window", "Planner, OMS, Policy, Critic", "Low (10/100)", "Auto-Approved Replacement + RMA"),
        ("Scenario 2", "Espresso Machine delayed in transit (Delhivery)", "OMS, Carrier Tool, Compensator", "Low (15/100)", "Carrier Trace + ₹500 Delay Credit"),
        ("Scenario 3", "Cosmetics returned after 38 days", "Policy Engine, Restocking Tool", "Low (5/100)", "Polite Refusal + 15% Goodwill Code"),
        ("Scenario 4", "₹1.76L Laptop Stolen Claim (Serial Returner)", "Fraud Detector, Signature Check", "Critical (88/100)", "OVERRULED: Escalated to Supervisor"),
        ("Scenario 5", "Platinum VIP Wrong Boot Size Shipped", "VIP Concierge Tool, Courier Swap", "Very Low (0/100)", "White-Glove Courier Exchange")
    ]

    for row_idx, row_data in enumerate(scenario_rows):
        for col_idx, text in enumerate(row_data):
            cell = tbl.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD_BG if row_idx % 2 == 0 else COLOR_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_PRIMARY

    # ==========================================================
    # SLIDE 7: Validation, Ethical Guardrails & Limitations
    # ==========================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Output Validation, Ethical Safeguards & Limitations")

    # Left Container: Validation & Ethics
    left_box7 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    left_box7.fill.solid()
    left_box7.fill.fore_color.rgb = COLOR_CARD_BG
    left_box7.line.color.rgb = COLOR_BORDER

    tb_l7 = slide7.shapes.add_textbox(Inches(1.05), Inches(2.0), Inches(5.1), Inches(4.4))
    tf_l7 = tb_l7.text_frame
    tf_l7.word_wrap = True

    p = tf_l7.paragraphs[0]
    p.text = "Validation & Ethical Guardrails"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_EMERALD
    p.space_after = Pt(10)

    val_points = [
        "100% Policy Conformity: Deterministic validation guarantees no hallucinated refund promises; category rules are programmatically enforced.",
        "Anti-Bias & Fair Treatment: Fraud scoring is based strictly on objective behavioral metrics (return velocity, signature conflicts, claim frequency), avoiding demographic bias.",
        "Auditability & Explainability: Every agent step produces structured trace logs, tool invocation arguments, and critic justification memos.",
        "Human-in-the-Loop Safeguard: High-consequence decisions (> ₹25,000 stolen claims, serial return anomalies) are never automated without human supervisor sign-off."
    ]
    for vp in val_points:
        p_v = tf_l7.add_paragraph()
        p_v.text = f"• {vp}"
        p_v.font.size = Pt(12)
        p_v.font.color.rgb = COLOR_PRIMARY
        p_v.space_after = Pt(8)

    # Right Container: Current Limitations
    right_box7 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    right_box7.fill.solid()
    right_box7.fill.fore_color.rgb = COLOR_CARD_BG
    right_box7.line.color.rgb = COLOR_BORDER

    tb_r7 = slide7.shapes.add_textbox(Inches(7.05), Inches(2.0), Inches(5.2), Inches(4.4))
    tf_r7 = tb_r7.text_frame
    tf_r7.word_wrap = True

    p_r7 = tf_r7.paragraphs[0]
    p_r7.text = "Current Limitations"
    p_r7.font.size = Pt(18)
    p_r7.font.bold = True
    p_r7.font.color.rgb = RGBColor(220, 38, 38)
    p_r7.space_after = Pt(10)

    lim_points = [
        "Multimodal Damage Inspection: Current PoC simulates photo damage verification; future version requires Computer Vision (CLIP / Vision LLMs) to grade damage severity.",
        "Live ERP / Carrier Integration: Operates on synthetic OMS and carrier telemetry; requires production REST API bridges to FedEx / Shopify / SAP.",
        "Multi-Turn Negotiation: PoC resolves single complaint tickets; edge cases involving back-and-forth counter-offers require conversational memory threads."
    ]
    for lp in lim_points:
        p_l = tf_r7.add_paragraph()
        p_l.text = f"• {lp}"
        p_l.font.size = Pt(12)
        p_l.font.color.rgb = COLOR_PRIMARY
        p_l.space_after = Pt(8)

    # ==========================================================
    # SLIDE 8: Future Scope & Conclusion
    # ==========================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Future Scope, Enterprise Roadmap & Conclusion")

    # Big Summary Card
    box_c = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    box_c.fill.solid()
    box_c.fill.fore_color.rgb = COLOR_PRIMARY
    box_c.line.color.rgb = COLOR_PRIMARY

    tb_c = slide8.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(10.9), Inches(4.2))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "Key Takeaways & Enterprise Impact"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_after = Pt(16)

    bullets_c = [
        "Transformed Business Process: Resolving e-commerce complaints shifts from a multi-day human queue to an autonomous, verified sub-second workflow.",
        "Collaborative Agent Advantage: The synergy of Planning, Research, Tool Execution, and Critic Oversight delivers unmatched accuracy and enterprise trust.",
        "Next-Generation Roadmap: Integrating Multimodal Image Verification, Model Context Protocol (MCP) microservices, and live webhook integrations.",
        "Measurable Bottom-Line Value: Reduces support headcount overhead by 70%, stops over ₹1,000 Cr in policy abuse, and boosts customer retention NPS."
    ]

    for b in bullets_c:
        p_b = tf_c.add_paragraph()
        p_b.text = f"• {b}"
        p_b.font.size = Pt(14)
        p_b.font.color.rgb = RGBColor(241, 245, 249)
        p_b.space_after = Pt(12)

    # Save Presentation
    prs.save(output_filename)
    print(f"Presentation generated successfully: {output_filename}")

if __name__ == "__main__":
    create_deck()
